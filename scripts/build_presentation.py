"""Питч-дек трека 1 (продажа + демо) → docs/presentation_track1.pptx"""

from __future__ import annotations

import os
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from make_demo_qr import build_qr  # noqa: E402

OUT = ROOT / "docs" / "presentation_track1.pptx"
OUT_FALLBACK = ROOT / "docs" / "presentation_track1_pitch.pptx"
QR_PATH = ROOT / "docs" / "demo_qr.png"

RED = RGBColor(0xE3, 0x06, 0x11)
RED_DARK = RGBColor(0xB3, 0x00, 0x0D)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5C, 0x5C, 0x62)
MUTED = RGBColor(0x8A, 0x8A, 0x90)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF2, 0xF2, 0xF4)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE4, 0xE4, 0xE8)

W = Inches(13.333)
H = Inches(7.5)


def _font(run, *, size=18, bold=False, color=DARK, name="Arial"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def _no_line(shape):
    shape.line.fill.background()


def _rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    _no_line(shape)
    return shape


def _round(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    _no_line(shape)
    # softer corners
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def _bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _p(tf, text, *, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT, space_after=6, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    _font(run, size=size, bold=bold, color=color)
    return p


def _set_text(shape, text, *, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set(qn("a:anchor"), "t" if anchor == MSO_ANCHOR.TOP else "ctr")
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _font(run, size=size, bold=bold, color=color)
    return tf


def _brand_bar(slide, label="Space CorpAcc with МТС · Трек 1"):
    _rect(slide, 0, 0, W, Inches(0.42), RED)
    box = _textbox(slide, Inches(0.55), Inches(0.06), Inches(10), Inches(0.32))
    tf = box.text_frame
    tf.word_wrap = False
    _p(tf, label, size=11, bold=True, color=WHITE, space_after=0, first=True)


def _footer(slide, page: str):
    box = _textbox(slide, Inches(0.55), Inches(7.1), Inches(10), Inches(0.28))
    tf = box.text_frame
    _p(tf, "ИИ-секретарь входящих · прототип для МТС", size=10, color=MUTED, space_after=0, first=True)
    num = _textbox(slide, Inches(12.2), Inches(7.1), Inches(0.8), Inches(0.28))
    _p(num.text_frame, page, size=10, color=MUTED, align=PP_ALIGN.RIGHT, space_after=0, first=True)


def _card(slide, left, top, width, height, title, body_lines, *, accent=True):
    card = _round(slide, left, top, width, height, CARD)
    if accent:
        _rect(slide, left, top, Inches(0.08), height, RED)
    tbox = _textbox(slide, left + Inches(0.28), top + Inches(0.22), width - Inches(0.45), Inches(0.4))
    _p(tbox.text_frame, title, size=15, bold=True, color=DARK, space_after=0, first=True)
    bbox = _textbox(
        slide,
        left + Inches(0.28),
        top + Inches(0.65),
        width - Inches(0.45),
        height - Inches(0.85),
    )
    tf = bbox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(body_lines):
        _p(tf, line, size=13, color=GRAY, space_after=4, first=(i == 0))
    return card


def _headline(slide, title, subtitle=None, top=Inches(0.65)):
    box = _textbox(slide, Inches(0.55), top, Inches(12.2), Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    _p(tf, title, size=30, bold=True, color=DARK, space_after=4, first=True)
    if subtitle:
        sbox = _textbox(slide, Inches(0.55), top + Inches(0.55), Inches(12.2), Inches(0.45))
        _p(sbox.text_frame, subtitle, size=15, color=GRAY, space_after=0, first=True)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    total = 13
    n = 0
    qr_file, demo_link = build_qr()

    def page():
        nonlocal n
        n += 1
        return f"{n}/{total}"

    # ——— 1. TITLE (sales opener) ———
    s = prs.slides.add_slide(blank)
    _bg(s, DARK)
    _rect(s, 0, 0, Inches(0.28), H, RED)
    _rect(s, 0, Inches(6.85), W, Inches(0.65), RED)
    box = _textbox(s, Inches(0.85), Inches(1.7), Inches(11.5), Inches(3.2))
    tf = box.text_frame
    tf.word_wrap = True
    _p(tf, "МТС · Трек 1", size=14, bold=True, color=RED, space_after=12, first=True)
    _p(tf, "ИИ-секретарь,", size=40, bold=True, color=WHITE, space_after=0)
    _p(tf, "который не пропускает важные звонки", size=40, bold=True, color=WHITE, space_after=14)
    _p(
        tf,
        "Услуга в приложении МТС + Telegram: подключение за минуты,\n"
        "голосовой ответ, резюме Ивану, контроль сценариев.",
        size=16,
        color=RGBColor(0xC8, 0xC8, 0xCE),
        space_after=0,
    )
    foot = _textbox(s, Inches(0.85), Inches(7.0), Inches(11), Inches(0.35))
    _p(
        foot.text_frame,
        "Питч + живое демо  ·  CJM: Иван Петров  ·  Space CorpAcc with МТС",
        size=12,
        bold=True,
        color=WHITE,
        space_after=0,
        first=True,
    )
    page()  # count title

    # ——— 2. PROBLEM ———
    s = prs.slides.add_slide(blank)
    _bg(s)
    _brand_bar(s)
    _headline(
        s,
        "Проблема: входящие съедают внимание и деньги",
        "Иван Петров, 42 — IT-предприниматель. Трубка звонит, когда он на совещании.",
    )
    _card(
        s,
        Inches(0.55),
        Inches(2.0),
        Inches(3.85),
        Inches(4.3),
        "Пропустил лид",
        [
            "Важный клиент не дозвонился.",
            "Сделка ушла конкуренту.",
            "Иван узнаёт об этом слишком поздно.",
        ],
    )
    _card(
        s,
        Inches(4.7),
        Inches(2.0),
        Inches(3.85),
        Inches(4.3),
        "Непрофессиональный ответ",
        [
            "Сам берёт трубку на бегу.",
            "Нет контекста, срывается тон.",
            "Бренд компании «плывёт».",
        ],
    )
    _card(
        s,
        Inches(8.85),
        Inches(2.0),
        Inches(3.85),
        Inches(4.3),
        "Рутина вместо бизнеса",
        [
            "Спам, FAQ, «ошибся номером».",
            "Часы в неделю — впустую.",
            "Нет силы нанять секретаря.",
        ],
    )
    _footer(s, page())

    # ——— 3. INSIGHT ———
    s = prs.slides.add_slide(blank)
    _bg(s)
    _brand_bar(s)
    _headline(s, "Инсайт", "Ивану нужен не «ещё один чат-бот», а секретарь на линии.")
    points = [
        ("≤ 5 минут", "Подключить без интеграторов и IT-отдела"),
        ("Контроль", "Свои сценарии, правила, правки ответов ИИ"),
        ("Эскалация", "Важное — сразу ему; рутину закрыть самому"),
        ("Канал МТС", "Уже в приложении и мессенджере, которым он пользуется"),
    ]
    for i, (t, b) in enumerate(points):
        top = Inches(1.9) + Inches(i * 1.15)
        _round(s, Inches(0.55), top, Inches(12.2), Inches(1.0), CARD)
        _rect(s, Inches(0.55), top, Inches(0.1), Inches(1.0), RED)
        tb = _textbox(s, Inches(0.95), top + Inches(0.18), Inches(2.8), Inches(0.65))
        _p(tb.text_frame, t, size=20, bold=True, color=RED, space_after=0, first=True)
        bb = _textbox(s, Inches(4.0), top + Inches(0.28), Inches(8.3), Inches(0.5))
        _p(bb.text_frame, b, size=16, color=DARK, space_after=0, first=True)
    _footer(s, page())

    # ——— 4. PRODUCT ———
    s = prs.slides.add_slide(blank)
    _bg(s)
    _brand_bar(s)
    _headline(
        s,
        "Что мы сделали",
        "Рабочий прототип услуги: приложение + бот + голосовой агент.",
    )
    _card(
        s,
        Inches(0.55),
        Inches(1.95),
        Inches(6.0),
        Inches(4.4),
        "Приложение МТС (Flet)",
        [
            "Вход по номеру → каталог → согласие → Подключить",
            "Шаблоны ответов: список, добавить, выключить",
            "Настройки: голос / текст / гибрид + правила",
            "История линии: вход, выход, резюме, правка",
            "Дашборд: важные / намерения / действия",
        ],
    )
    _card(
        s,
        Inches(6.85),
        Inches(1.95),
        Inches(5.9),
        Inches(4.4),
        "Telegram + голос",
        [
            "Кнопка бота сразу после подключения",
            "Демо-звонок: приветствие голосом",
            "ГС абонента → STT → LLM → TTS → ответ",
            "Пуши и дашборд Ивану в мессенджере",
            "Карточка звонка попадает в историю",
        ],
    )
    _footer(s, page())

    # ——— 5. VALUE ———
    s = prs.slides.add_slide(blank)
    _bg(s)
    _brand_bar(s)
    _headline(s, "Ценность для Ивана — и для МТС")
    _card(
        s,
        Inches(0.55),
        Inches(1.7),
        Inches(6.0),
        Inches(4.7),
        "Для Ивана",
        [
            "Не пропускает важное, пока на совещании",
            "Профессиональный голос бренда на линии",
            "Резюме и расшифровка — в один клик",
            "Правит сценарии без разработчика",
            "Горячая линия (заглушка демо) — если ИИ ошибся",
        ],
    )
    _card(
        s,
        Inches(6.85),
        Inches(1.7),
        Inches(5.9),
        Inches(4.7),
        "Для МТС",
        [
            "Новая услуга в уже существующих каналах",
            "Дистрибуция: приложение + бот",
            "Потенциал ARPU малого бизнеса",
            "Сценарий монетизации: с баланса / пакеты",
            "Путь в SIP / ЦОД / биллинг оператора",
        ],
    )
    _footer(s, page())

    # ——— 6. HOW WE BUILT ———
    s = prs.slides.add_slide(blank)
    _bg(s)
    _brand_bar(s)
    _headline(s, "Как сделали", "Стек под демо жюри: быстро поднять, честно показать fallback.")
    layers = [
        ("Клиент", "Flet-приложение «как МТС»  ·  Telegram-бот (aiogram)"),
        ("API", "FastAPI  ·  история / настройки / правила / сценарии / hotline"),
        ("Интеллект", "YandexGPT + routing rules  ·  training examples в промпте"),
        ("Речь", "SpeechKit STT/TTS  ·  fallback: Whisper + pyttsx3"),
        ("Данные", "SQLite по линии  ·  секреты только в .env (не в git)"),
    ]
    for i, (k, v) in enumerate(layers):
        y = Inches(1.75) + Inches(i * 0.9)
        _round(s, Inches(0.55), y, Inches(12.2), Inches(0.78), CARD)
        kbox = _textbox(s, Inches(0.85), y + Inches(0.2), Inches(2.2), Inches(0.4))
        _p(kbox.text_frame, k, size=15, bold=True, color=RED, space_after=0, first=True)
        vbox = _textbox(s, Inches(3.3), y + Inches(0.2), Inches(9.0), Inches(0.4))
        _p(vbox.text_frame, v, size=15, color=DARK, space_after=0, first=True)
    _footer(s, page())

    # ——— 7. TZ FIT (compact, sales not checklist dump) ———
    s = prs.slides.add_slide(blank)
    _bg(s)
    _brand_bar(s)
    _headline(s, "Закрываем ТЗ трека 1", "Не чеклист ради галочек — то, что Иван реально получит.")
    rows = [
        ("Подключение ≤ 5 мин", "Онбординг в приложении, /ready"),
        ("Голос или чат", "Способ ответа + action_required агента"),
        ("Суть обращения", "STT + резюме + карточка в истории"),
        ("Контроль Ивана", "Шаблоны, правила, правка ответа ИИ"),
        ("Конфиденциальность", "Согласие, disclosure, секреты, fallback локально"),
    ]
    for i, (a, b) in enumerate(rows):
        y = Inches(1.75) + Inches(i * 0.9)
        _round(s, Inches(0.55), y, Inches(12.2), Inches(0.78), CARD)
        _rect(s, Inches(0.55), y, Inches(0.1), Inches(0.78), RED)
        ab = _textbox(s, Inches(0.9), y + Inches(0.2), Inches(4.5), Inches(0.4))
        _p(ab.text_frame, a, size=16, bold=True, color=DARK, space_after=0, first=True)
        bb = _textbox(s, Inches(5.6), y + Inches(0.2), Inches(6.8), Inches(0.4))
        _p(bb.text_frame, b, size=15, color=GRAY, space_after=0, first=True)
    _footer(s, page())

    # ——— 8. DEMO SCRIPT (centerpiece) ———
    s = prs.slides.add_slide(blank)
    _bg(s)
    _brand_bar(s, "Живое демо · ~3–4 минуты")
    _headline(
        s,
        "Сценарий демонстрации",
        "Показываем путь Ивана целиком — от подключения до карточки звонка.",
    )
    steps = [
        ("01", "Приложение", "Вход по номеру → Подключить услугу"),
        ("02", "Контроль", "Шаблон приветствия + правило (человек / спам)"),
        ("03", "Бот", "Открыть Telegram → «Начать звонок»"),
        ("04", "Голос", "Агент здоровается голосом → вы отвечаете ГС"),
        ("05", "Итог", "Ответ голосом + резюме в истории / важных"),
    ]
    for i, (num, title, desc) in enumerate(steps):
        x = Inches(0.45) + Inches(i * 2.55)
        _round(s, x, Inches(2.1), Inches(2.4), Inches(3.9), CARD)
        nb = _textbox(s, x + Inches(0.2), Inches(2.35), Inches(2.0), Inches(0.45))
        _p(nb.text_frame, num, size=22, bold=True, color=RED, space_after=0, first=True)
        tb = _textbox(s, x + Inches(0.2), Inches(2.95), Inches(2.0), Inches(0.45))
        _p(tb.text_frame, title, size=16, bold=True, color=DARK, space_after=0, first=True)
        db = _textbox(s, x + Inches(0.2), Inches(3.55), Inches(2.0), Inches(2.0))
        _p(db.text_frame, desc, size=13, color=GRAY, space_after=0, first=True)
    tip = _textbox(s, Inches(0.55), Inches(6.25), Inches(12.2), Inches(0.55))
    _p(
        tip.text_frame,
        "Запасной план: если микрофон капризничает — текстовая реплика в боте. Эскалация на человека — честная заглушка демо.",
        size=13,
        color=GRAY,
        space_after=0,
        first=True,
    )
    _footer(s, page())

    # ——— 9. DEMO MOMENTS (what jury should feel) ———
    s = prs.slides.add_slide(blank)
    _bg(s)
    _brand_bar(s)
    _headline(s, "Что жюри должно увидеть в демо", "Три «вау»-момента — не технический тур.")
    moments = [
        (
            "«Подняли за минуты»",
            "Из каталога услуг — в рабочую линию без кода и SIP-настроек.",
        ),
        (
            "«Говорит как секретарь»",
            "Голосовое приветствие и ответ по сути; Ивану — короткое резюме.",
        ),
        (
            "«Иван управляет»",
            "Поменяли шаблон / правило — и это видно в поведении на следующем звонке.",
        ),
    ]
    for i, (t, b) in enumerate(moments):
        y = Inches(1.9) + Inches(i * 1.45)
        _round(s, Inches(0.55), y, Inches(12.2), Inches(1.25), CARD)
        _rect(s, Inches(0.55), y, Inches(0.12), Inches(1.25), RED)
        tb = _textbox(s, Inches(1.0), y + Inches(0.22), Inches(11.3), Inches(0.4))
        _p(tb.text_frame, t, size=20, bold=True, color=DARK, space_after=0, first=True)
        bb = _textbox(s, Inches(1.0), y + Inches(0.65), Inches(11.3), Inches(0.4))
        _p(bb.text_frame, b, size=15, color=GRAY, space_after=0, first=True)
    _footer(s, page())

    # ——— 10. INTEGRATION / MONEY ———
    s = prs.slides.add_slide(blank)
    _bg(s)
    _brand_bar(s)
    _headline(
        s,
        "Как встраиваемся в МТС",
        "Прототип уже сидит в каналах оператора. Дальше — инфраструктура МТС, не наш pet-проект.",
    )
    _card(
        s,
        Inches(0.55),
        Inches(1.85),
        Inches(4.0),
        Inches(4.5),
        "Сейчас (демо)",
        [
            "UX приложения МТС",
            "Бот как канал пушей",
            "YandexGPT + SpeechKit",
            "Локальный fallback",
        ],
    )
    _card(
        s,
        Inches(4.75),
        Inches(1.85),
        Inches(4.0),
        Inches(4.5),
        "Следующий шаг",
        [
            "SIP / стрим в трубку",
            "Биллинг с баланса",
            "МТС ID / auth",
            "ЦОД МТС, без зарубежного облака для ПДн",
        ],
    )
    _card(
        s,
        Inches(8.95),
        Inches(1.85),
        Inches(3.8),
        Inches(4.5),
        "Монетизация",
        [
            "Услуга в каталоге",
            "Пакеты Basic / Plus",
            "ARPU micro-business",
            "Потом B2B-линии",
        ],
    )
    _footer(s, page())

    # ——— 11. TRUST ———
    s = prs.slides.add_slide(blank)
    _bg(s)
    _brand_bar(s)
    _headline(s, "Доверие на линии", "Звонок — чувствительные данные. Это часть продажи.")
    items = [
        ("Согласие", "Оферта при подключении услуги"),
        ("Disclosure", "Абоненту ясно: отвечает ИИ"),
        ("Изоляция", "История только своей линии"),
        ("Секреты", "Ключи в .env, не в репозитории"),
        ("Прод", "Контур МТС / РФ — roadmap на защите"),
    ]
    for i, (t, b) in enumerate(items):
        x = Inches(0.45) + Inches((i % 5) * 2.55)
        _round(s, x, Inches(2.2), Inches(2.4), Inches(3.5), CARD)
        _rect(s, x, Inches(2.2), Inches(2.4), Inches(0.12), RED)
        tb = _textbox(s, x + Inches(0.18), Inches(2.6), Inches(2.05), Inches(0.8))
        _p(tb.text_frame, t, size=16, bold=True, color=DARK, space_after=0, first=True)
        bb = _textbox(s, x + Inches(0.18), Inches(3.5), Inches(2.05), Inches(1.8))
        _p(bb.text_frame, b, size=13, color=GRAY, space_after=0, first=True)
    _footer(s, page())

    # ——— 12. TRY IT (QR = Flet / МТС веб, не бот) ———
    s = prs.slides.add_slide(blank)
    _bg(s)
    _brand_bar(s, "Попробуйте сами · приложение МТС")
    _headline(
        s,
        "Откройте веб-приложение по QR",
        "Лёгкая страница МТС на телефоне. Telegram — только кнопка с личным ключом.",
    )
    _round(s, Inches(0.55), Inches(1.85), Inches(6.4), Inches(4.6), CARD)
    left = _textbox(s, Inches(0.85), Inches(2.15), Inches(5.8), Inches(4.0))
    tf = left.text_frame
    tf.word_wrap = True
    _p(tf, "Как зайти", size=18, bold=True, color=DARK, space_after=10, first=True)
    lines = (
        "1. Сканируйте QR → сайт МТС",
        "2. Номер → согласие → Продолжить",
        "3. «Перейти в Telegram» — личный ключ",
        "4. Подтвердить → демо-звонок в боте",
        "",
        f"Ссылка: {demo_link}",
        "",
        "Лёгкий веб (~8KB), не Flet 10MB.",
        "Перед защитой: tunnel на :8000.",
    )
    for line in lines:
        muted = line.startswith("QR ") or line.startswith("Жюри")
        _p(tf, line, size=14, color=GRAY if muted else DARK, space_after=6)
    _round(s, Inches(7.3), Inches(1.85), Inches(5.3), Inches(4.6), CARD)
    if qr_file.is_file():
        s.shapes.add_picture(
            str(qr_file),
            Inches(8.15),
            Inches(2.25),
            width=Inches(3.6),
            height=Inches(3.6),
        )
    url_box = _textbox(s, Inches(7.5), Inches(5.95), Inches(4.9), Inches(0.35))
    _p(url_box.text_frame, demo_link, size=11, color=MUTED, align=PP_ALIGN.CENTER, space_after=0, first=True)
    _footer(s, page())

    # ——— 13. CLOSE + ASK ———
    s = prs.slides.add_slide(blank)
    _bg(s, DARK)
    _rect(s, 0, 0, Inches(0.28), H, RED)
    box = _textbox(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(3.5))
    tf = box.text_frame
    tf.word_wrap = True
    _p(tf, "Итог", size=14, bold=True, color=RED, space_after=10, first=True)
    _p(tf, "Иван не пропускает важное.", size=32, bold=True, color=WHITE, space_after=6)
    _p(tf, "МТС получает услугу в своих каналах.", size=32, bold=True, color=WHITE, space_after=18)
    _p(
        tf,
        "Дальше — живое демо и QR. Вопросы после.",
        size=18,
        color=RGBColor(0xC8, 0xC8, 0xCE),
        space_after=0,
    )
    _rect(s, 0, Inches(6.85), W, Inches(0.65), RED)
    ask = _textbox(s, Inches(0.9), Inches(7.0), Inches(11.5), Inches(0.35))
    _p(
        ask.text_frame,
        "Ask: поддержка пилота в каналах МТС  ·  обратная связь по CJM Ивана",
        size=13,
        bold=True,
        color=WHITE,
        space_after=0,
        first=True,
    )
    page()

    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(OUT)
        return OUT
    except PermissionError:
        prs.save(OUT_FALLBACK)
        return OUT_FALLBACK


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path} ({path.stat().st_size} bytes)")
