"""Insert customer slide «Тесты и аналитика» as #7 into the current 8-slide deck."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from build_presentation import (  # noqa: E402
    MUTED,
    RED,
    WHITE,
    W,
    _bg,
    _card,
    _footer,
    _headline,
    _p,
    _rect,
    _textbox,
)

PPTX = ROOT / "docs" / "presentation_track1.pptx"
FALLBACK = ROOT / "docs" / "presentation_track1_pitch.pptx"


def main() -> Path:
    prs = Presentation(str(PPTX))
    # Already inserted?
    for slide in prs.slides:
        for sh in slide.shapes:
            if hasattr(sh, "text") and "Тесты и аналитика" in (sh.text or ""):
                print("Slide already present; renumber only")
                _renumber(prs)
                return _save(prs)

    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    _bg(s)
    _rect(s, 0, 0, W, Inches(0.42), RED)
    brand = _textbox(s, Inches(0.55), Inches(0.06), Inches(10), Inches(0.32))
    _p(
        brand.text_frame,
        "StartUp Space VSTU",
        size=11,
        bold=True,
        color=WHITE,
        space_after=0,
        first=True,
    )
    _headline(
        s,
        "Тесты и аналитика",
        "Заказчику видно: качество под контролем, а Иван видит картину по линии.",
    )
    _card(
        s,
        Inches(0.55),
        Inches(1.85),
        Inches(6.0),
        Inches(4.5),
        "Тесты (качество)",
        [
            "Контракты API и матрица маршрутизации",
            "Живой звонок: VAD, barge-in, многоходовый диалог",
            "UI приложения и дашборд Telegram",
            "Fallback: если Yandex недоступен — Whisper / edge-tts",
            "Изоляция данных: история только своей линии",
        ],
    )
    _card(
        s,
        Inches(6.85),
        Inches(1.85),
        Inches(5.9),
        Inches(4.5),
        "Аналитика (ценность Ивану)",
        [
            "Дашборд: важные / рутина / доля critical",
            "История звонков + фильтр + правка карточки",
            "Статистика по intent / action / priority",
            "Пуш в Telegram с резюме после звонка",
            "Метрики линии — без Excel и ручного разбора",
        ],
    )
    _footer(s, "7/9")

    sld_id_lst = prs.slides._sldIdLst
    ids = list(sld_id_lst)
    last = ids[-1]
    sld_id_lst.remove(last)
    sld_id_lst.insert(6, last)

    _renumber(prs)
    return _save(prs)


def _renumber(prs: Presentation) -> None:
    pat = re.compile(r"^\d+/\d+$")
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not hasattr(sh, "text") or not sh.has_text_frame:
                continue
            t = (sh.text or "").strip()
            if not pat.match(t):
                continue
            sh.text_frame.clear()
            p = sh.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT
            run = p.add_run()
            run.text = f"{i}/{total}"
            run.font.size = Pt(10)
            run.font.color.rgb = MUTED
            run.font.name = "Arial"


def _save(prs: Presentation) -> Path:
    try:
        prs.save(str(PPTX))
        return PPTX
    except PermissionError:
        prs.save(str(FALLBACK))
        return FALLBACK


if __name__ == "__main__":
    path = main()
    prs = Presentation(str(path))
    print(f"Wrote {path} ({len(prs.slides)} slides)")
    for i, slide in enumerate(prs.slides, 1):
        titles = [
            sh.text.strip().split("\n")[0][:70]
            for sh in slide.shapes
            if hasattr(sh, "text") and sh.text.strip()
        ]
        label = titles[1] if len(titles) > 1 else (titles[0] if titles else "?")
        print(f"  {i}. {label}")
